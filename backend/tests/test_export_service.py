"""Tests for PDF executive dossiers and Excel batch workbooks export engine."""

import io
import openpyxl
from starlette.testclient import TestClient

from backend.app.main import app
from backend.app.services.export_service import (
    generate_batch_excel_portfolio,
    generate_cuit_pdf_dossier,
)


def test_generate_cuit_pdf_dossier_structure():
    """Verify generated PDF dossier has valid binary structure and header."""
    mock_debtor = {
        "cuit": "20-30123456-3",
        "denominacion": "ACME S.A.",
        "worst_situation": 2,
        "total_debt": 150000.0,
        "periodo": "2024-05",
        "alert_badges": [
            {
                "id": "CREDIT_DETERIORATION",
                "name": "Deterioro Crediticio",
                "severity": "high",
                "description": "Empeoramiento de situacion",
            }
        ],
        "entidades": [
            {"entidad": "BANCO DE GALICIA", "situacion": 1, "monto": 100.5},
            {"entidad": "BANCO SANTANDER", "situacion": 2, "monto": 49.5},
        ],
        "cheques_rechazados": {"cantidad": 2, "total_monto": 25000.0},
    }

    pdf_bytes = generate_cuit_pdf_dossier(mock_debtor)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 200
    assert pdf_bytes.startswith(b"%PDF-1.4")
    assert b"%%EOF" in pdf_bytes


def test_generate_batch_excel_portfolio_structure():
    """Verify generated Excel workbook has valid structure, worksheets, and headers."""
    batch_data = {
        "metrics": {
            "total_cuits": 2,
            "valid_cuits": 2,
            "invalid_cuits": 0,
            "max_situacion": 2,
            "total_alerts": 1,
        },
        "results": [
            {
                "cuit": "20-30123456-3",
                "denominacion": "EMPRESA ALFA S.A.",
                "worst_situation": 1,
                "total_debt": 50000.0,
                "entidades": [{"entidad": "BANCO NACION"}],
                "cheques_rechazados": {"cantidad": 0, "total_monto": 0.0},
                "alert_badges": [],
            },
            {
                "cuit": "30-71234567-1",
                "denominacion": "EMPRESA BETA S.R.L.",
                "worst_situation": 2,
                "total_debt": 98000.0,
                "entidades": [{"entidad": "BANCO SANTANDER"}],
                "cheques_rechazados": {"cantidad": 1, "total_monto": 12000.0},
                "alert_badges": [{"name": "Alerta Cheques"}],
            },
        ],
    }

    excel_bytes = generate_batch_excel_portfolio(batch_data)
    assert isinstance(excel_bytes, bytes)
    assert excel_bytes.startswith(b"PK")  # ZIP container signature for xlsx

    # Load with openpyxl to verify spreadsheet validity
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Resumen Portafolio" in wb.sheetnames
    ws = wb["Resumen Portafolio"]
    assert "INFORME CONSOLIDADO" in str(ws["A1"].value)
    assert ws.max_row >= 8  # Banner + metrics + headers + 2 data rows


def test_export_pdf_endpoint_success(client):
    """Test GET /api/export/cuit-dossier/{cuit} returns 200 OK with PDF attachment."""
    # Valid test CUIT from mock data
    cuit = "20-30123456-3"
    response = client.get(f"/api/export/cuit-dossier/{cuit}")
    assert response.status_code == 200
    assert "application/pdf" in response.headers.get("content-type", "")
    assert "attachment" in response.headers.get("content-disposition", "")
    assert response.content.startswith(b"%PDF-1.4")


def test_export_pdf_endpoint_invalid_cuit(client):
    """Test GET /api/export/cuit-dossier/{cuit} with invalid CUIT returns 400 Bad Request."""
    response = client.get("/api/export/cuit-dossier/11-11111111-1")
    assert response.status_code == 400
    assert "inválido" in response.json()["detail"].lower()


def test_export_excel_endpoint_with_cuits_list(client):
    """Test POST /api/export/batch-excel with list of CUITs returns 200 OK with xlsx file."""
    payload = {"cuits": ["20-30123456-3", "30-71234567-1"]}
    response = client.post("/api/export/batch-excel", json=payload)
    assert response.status_code == 200
    assert (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        in response.headers.get("content-type", "")
    )
    assert "attachment" in response.headers.get("content-disposition", "")
    assert response.content.startswith(b"PK")


def test_export_excel_endpoint_with_precomputed_results(client):
    """Test POST /api/export/batch-excel with pre-evaluated results returns 200 OK."""
    payload = {
        "results": [
            {
                "cuit": "20-30123456-3",
                "denominacion": "EMPRESA DEMO",
                "worst_situation": 1,
                "total_debt": 10000.0,
                "entidades": [],
                "cheques_rechazados": {"cantidad": 0, "total_monto": 0.0},
                "alert_badges": [],
            }
        ]
    }
    response = client.post("/api/export/batch-excel", json=payload)
    assert response.status_code == 200
    assert response.content.startswith(b"PK")


def test_export_excel_endpoint_empty_request_returns_400(client):
    """Test POST /api/export/batch-excel with empty payload returns 400 Bad Request."""
    response = client.post("/api/export/batch-excel", json={})
    assert response.status_code == 400
    assert "Debe proveer" in response.json()["detail"]
